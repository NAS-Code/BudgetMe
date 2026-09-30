"""
Rebuilds the labeled history from the original monthly tracker workbook.

  python scripts/extract_history.py [path/to/Monthly_Budget_Tracker.xlsx]

Writes data/labeled_transactions.csv, data/merchant_lookup.csv and
data/budget_history_labeled.xlsx. Skips the TEMPLATE and Planned Budget
Increase tabs; reads columns B-F of every monthly tab.
"""
from openpyxl import load_workbook
import datetime as dt, re, pandas as pd
import os, sys
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WB=sys.argv[1] if len(sys.argv)>1 else os.path.join(ROOT,'data','source','Monthly_Budget_Tracker.xlsx')
wb=load_workbook(WB,read_only=True,data_only=True)
VAR={'Week 1 Fun':('Fun',1),'Week 2 Fun':('Fun',2),'Week 3 Fun':('Fun',3),'Week 4 Fun':('Fun',4),'Week 5 Fun':('Fun',5),
     'Transportation':('Transportation',None),'Grocery/Food':('Grocery/Food',None),
     'MISCELLANEOUS EXPENSES':('Miscellaneous',None),'Work Expenses':('Work Expenses',None)}
CARRY={'Monthly Surplus','Extra Leftover','Leftover'}
SAVE={'Savings','Crypto','Travel Fund','Danger Fund'}
num=lambda v: isinstance(v,(int,float)) and not isinstance(v,bool)
def key(desc):
    s=desc.lower().strip()
    s=re.sub(r"[’'`]", "", s); s=re.sub(r"\s+"," ",s)
    return s
rows=[]; unequal=[]
for name in wb.sheetnames[2:]:
    m=dt.datetime.strptime(name.title(),'%B %Y')
    ws=wb[name]; sec=None; colE=None
    for r,row in enumerate(ws.iter_rows(max_col=6,values_only=True),1):
        b,c,d,e,f=row[1],row[2],row[3],row[4],row[5]
        bs=b.strip() if isinstance(b,str) else b
        if bs=='INCOME': sec='Income'; continue
        if bs=='MONTHLY EXPENSES': sec='Monthly'; continue
        if bs=='VARIABLE EXPENSES': sec=None; continue
        if bs in VAR and c=='DATE': sec=VAR[bs]; colE=str(e); continue
        if bs=='DESCRIPTION': continue
        if d=='TOTAL': sec=None; continue
        if sec is None or not isinstance(bs,str) or not bs: continue
        date=c.date() if isinstance(c,dt.datetime) else None
        base=dict(month=m.strftime('%Y-%m'),date=date,description=bs,source=f"'{name}'!B{r}")
        if sec=='Income':
            if bs in CARRY or not num(f) or f==0: continue
            rows.append({**base,'kind':'income','category':'Income','subcategory':bs,'fun_week':None,'account':'Wells Fargo','amount':f})
        elif sec=='Monthly':
            if not num(f) or f==0: continue
            k='savings_transfer' if bs in SAVE else 'fixed'
            rows.append({**base,'kind':k,'category':'Monthly','subcategory':bs,'fun_week':None,'account':'Wells Fargo','amount':f})
        else:
            cat,wk=sec
            cc = colE=='Credit Card' and num(e) and e!=0
            wf = num(f) and f!=0
            if cc:
                rows.append({**base,'kind':'variable','category':cat,'subcategory':None,'fun_week':wk,'account':'Chase Credit Card','amount':e})
                if wf and abs(e-f)>0.005: unequal.append((name,r,bs,e,f))
            elif wf:
                rows.append({**base,'kind':'variable','category':cat,'subcategory':None,'fun_week':wk,'account':'Wells Fargo','amount':f})
df=pd.DataFrame(rows)
df['merchant_key']=df.description.map(key)
# Venmo notes
v=df.description.str.extract(r'(?i)^venmo\s*\((.*)\)\s*$')[0]
df['note']=v
df.loc[v.notna(),'merchant_key']='venmo'
df['is_refund']=df.amount<0
flags=[]
for i,x in df.iterrows():
    fl=[]
    if x.account=='Chase Credit Card' and x.merchant_key=='subway' and x.amount==150 and x.date is None: fl.append('likely template placeholder')
    if x.date is not None and x.date.strftime('%Y-%m')!=x.month: fl.append('date outside sheet month')
    flags.append('; '.join(fl))
df['flags']=flags
dup=df[df.kind=='variable'].duplicated(['month','date','merchant_key','amount'],keep=False) & df.date.notna()
df.loc[dup[dup].index,'flags']=df.loc[dup[dup].index,'flags'].map(lambda s:(s+'; ' if s else '')+'same charge in two sections')
cols=['month','date','description','merchant_key','note','amount','is_refund','account','kind','category','subcategory','fun_week','flags','source']
df=df[cols].sort_values(['month','date','source'],na_position='last').reset_index(drop=True)

# ---- normalize merchants, flags, lookup ----
import unicodedata
# parentheticals -> note
par=df.description.str.extract(r'^(.*?)\s*\((.*)\)\s*$')
has=par[0].notna() & (par[0].str.strip()!='')
df.loc[has,'note']=par.loc[has,1]
base=df.description.where(~has, par[0])
GEN={'deli','market','bodega','bagel','bagels','grocery','cafe','restaurant','inc','nyc','the','pharmacy'}
def grp(s):
    s=unicodedata.normalize('NFKD',s.lower()); s=''.join(ch for ch in s if not unicodedata.combining(ch))
    s=s.replace('&',' and '); s=re.sub(r"[^a-z0-9 ]"," ",s)
    w=[t for t in s.split() if t]
    w2=[t for t in w if t not in GEN] or w
    k=''.join(w2)
    if len(k)>4 and k.endswith('s'): k=k[:-1]
    return k
df['merchant']=base.str.strip()
df['merchant_key']=df.merchant.map(grp)
isv=df.merchant_key=='venmo'
df.loc[df.kind.eq('variable') & df.description.str.match(r'(?i)^(difference |crazy )?adjustment$'),'kind']='adjustment'
# recompute cross-section dup flag
def fixflag(s): return '; '.join(p for p in s.split('; ') if p and p!='same charge in two sections')
df['flags']=df['flags'].map(fixflag)
v=df[(df.kind=='variable')&df.date.notna()]
g=v.groupby(['month','date','merchant_key','amount']).category.nunique()
cross=g[g>1].index
m=df.set_index(['month','date','merchant_key','amount']).index.isin(cross) & (df.kind=='variable')
df.loc[m,'flags']=df.loc[m,'flags'].map(lambda s:(s+'; ' if s else '')+'same charge filed in two categories')
cols=['month','date','merchant','merchant_key','note','description','amount','is_refund','account','kind','category','subcategory','fun_week','flags','source']
df=df[cols]
# merchant lookup
p=df[df.kind=='variable']
rows=[]
for k,x in p.groupby('merchant_key'):
    vc=x.category.value_counts()
    rows.append(dict(merchant_key=k,
        example_names=', '.join(x.merchant.value_counts().index[:3]),
        times_seen=len(x), top_category=vc.index[0], top_category_share=round(vc.iloc[0]/len(x),3),
        category_breakdown=', '.join(f'{c} {n}' for c,n in vc.items()),
        median_amount=round(x.amount.median(),2), first_seen=x.month.min(), last_seen=x.month.max(),
        accounts=', '.join(sorted(x.account.unique()))))
lk=pd.DataFrame(rows).sort_values(['times_seen','merchant_key'],ascending=[False,True])

# ---- write outputs ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
out=os.path.join(ROOT,'data')+os.sep
df.to_csv(out+'labeled_transactions.csv',index=False)
lk.to_csv(out+'merchant_lookup.csv',index=False)
wb=Workbook()
F=Font(name='Arial',size=10); H=Font(name='Arial',size=10,bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='305496')
def sheet(ws,d,widths,money=(),dates=(),pct=()):
    ws.append(list(d.columns))
    for r in d.itertuples(index=False): ws.append([None if (isinstance(v,float) and pd.isna(v)) else v for v in r])
    for c in ws[1]: c.font=H; c.fill=fill; c.alignment=Alignment(vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=F
    for i,col in enumerate(d.columns,1):
        L=get_column_letter(i); ws.column_dimensions[L].width=widths.get(col,14)
        fmt='$#,##0.00;($#,##0.00);-' if col in money else 'yyyy-mm-dd' if col in dates else '0%' if col in pct else None
        if fmt:
            for c in ws[L][1:]: c.number_format=fmt
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
ws=wb.active; ws.title='Transactions'
sheet(ws,df,{'merchant':26,'merchant_key':18,'note':28,'description':34,'account':18,'kind':16,'category':16,'subcategory':24,'flags':30,'source':22,'fun_week':9,'is_refund':9},money={'amount'},dates={'date'})
ws=wb.create_sheet('Merchant Lookup')
sheet(ws,lk,{'merchant_key':20,'example_names':38,'top_category':16,'category_breakdown':44,'accounts':30,'top_category_share':12},money={'median_amount'},pct={'top_category_share'})
ws=wb.create_sheet('Notes')
notes=[('About this file',''),
('Source','Generated from Monthly_Budget_Tracker.xlsx, every month from May 2021 to October 2026. The TEMPLATE and Planned Budget Increase tabs are skipped. Values are a static snapshot, not formulas.'),
('source column','Points back to the exact cell in your original workbook, so any row can be checked.'),
('kind','variable = purchases in your Fun / Transportation / Grocery/Food / Misc / Work sections. fixed = Monthly Expenses line items. savings_transfer = Savings, Crypto, Travel Fund, Danger Fund. income = income lines (Monthly Surplus / Leftover carryovers excluded). adjustment = your balancing entries, kept out of the merchant lookup.'),
('account','Chase Credit Card where you logged the amount in the Credit Card column (only used from mid-2026). Everything else is Wells Fargo, the only column before that.'),
('fun_week','For Fun purchases, which weekly bucket you filed it in (1 to 5).'),
('merchant / note','Text in parentheses was split out as a note, e.g. "Venmo (Attaboy)" becomes merchant Venmo, note Attaboy. That note is the intent signal a bank feed will not have.'),
('merchant_key','Normalized name used for matching: lowercase, punctuation and generic words (deli, market, bodega, pharmacy...) removed, trailing s dropped. Groups "Key Food" with "Key Foods", "Food Town" with "Foodtown".'),
('flags','likely template placeholder = the $150 Subway line carried over from the template. date outside sheet month = a typo or a late entry. same charge filed in two categories = same merchant, date and amount under two headings.'),
('Merchant Lookup','One row per merchant_key across variable purchases: how often it appears, the category you chose most, and how consistently.'),
('Not included','Amounts in ESTIMATE columns (plans, not spending), TOTAL rows, and anything outside columns B to F.')]
for a,b in notes: ws.append([a,b])
for row in ws.iter_rows():
    for c in row: c.font=F; c.alignment=Alignment(wrap_text=True,vertical='top')
ws['A1'].font=Font(name='Arial',size=12,bold=True)
ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=100
wb.save(out+'budget_history_labeled.xlsx')
print('ok')

print(len(df),'transactions,',len(lk),'merchants ->',out)
